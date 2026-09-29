"""Atualiza a nota de decisoes metodologicas com alteracoes controladas, alinhando-a ao artigo revisado.
Uso: python aplicar_nota.py valores.json  ->  artigo/Nota_Decisoes_Metodologicas_revisada.docx"""
import sys, json, os
import docx
from tracked import replace, insert_after, consertar_midia

AQUI = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(AQUI, '..', 'artigo', 'Nota_Decisoes_Metodologicas.docx')
DST = os.path.join(AQUI, '..', 'artigo', 'Nota_Decisoes_Metodologicas_revisada.docx')
V = json.load(open(sys.argv[1]))
log = []
consertar_midia(SRC, DST)
d = docx.Document(DST)


def paragrafos():
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


def sub(old, new, motivo):
    achados = []
    for p in paragrafos():
        if old in p.text and all(p._p is not q._p for q in achados):
            achados.append(p)
    if len(achados) != 1:
        log.append(f'FALHOU ({len(achados)}): {old[:70]!r}'); return
    replace(achados[0], old, new)
    log.append(f'OK     {motivo}')


# secao 2
sub('a única diferença significativa é contra o escore supervisionado, que ele supera por 0,066 ponto.',
    'a única diferença significativa é contra o escore supervisionado, que ele supera por 0,066 ponto, resultado que não resiste à antecipação das datas de recuperação judicial descrita na seção 6.',
    'secao 2: resultado fragil')
sub('ficam em 0,857, 0,856 e 0,858.', 'ficam em 0,857, 0,856 e 0,857.', 'secao 2: arvores')
sub('A explicação é de composição e foi verificada nos dados.',
    'No horizonte de um ano, a maior parte do contraste decorre da sobreposição mecânica: as 66 entradas têm EBITDA negativo já no exercício da previsão, e a margem EBITDA isolada alcança 0,978 nesse subconjunto. No horizonte de dois anos, em que a sobreposição desaparece, os dados são consistentes com uma explicação de composição.',
    'secao 2: mecanico e composicao')
sub('e sim cego a uma das rotas de deterioração', 'e sim fraco em uma das rotas de deterioração', 'secao 2: moderacao')

# quadro 1
sub('área sob a curva de 0,783 para 0,770', f"área sob a curva de {V['auc8']} para {V['auc7']}", 'I.2: Ponzi com exercicios consecutivos')
sub('e excluí-las praticamente não altera nenhum resultado.',
    'e excluí-las praticamente não altera nenhum resultado. Antecipar em um ano a data das 22 firmas cujo primeiro documento trata do andamento do processo também não altera a conclusão central.',
    'I.3: teste da RJ atrasada')
sub('Das 143 firmas que saem, 32,2% estavam em vulnerabilidade na última observação, contra 16,0% das que permanecem; 46 saídas têm sinal de dificuldade financeira e 6 ocorrem a partir do estado saudável, tratadas como evento na sensibilidade.',
    f"Das {V['n_saem']} firmas que saem, {V['vul_saem']} estavam em vulnerabilidade na última observação, contra {V['vul_ficam']} das que permanecem. Pelo cadastro, as saídas dividem-se em fechamento voluntário ({V['c_fech']}), recuperação judicial ({V['c_rj']}), incorporação ({V['c_inc']}), falência ou liquidação ({V['c_fal']}), outros cancelamentos ({V['c_out']}) e motivo não identificado ({V['c_nid']}); tratar como evento as saídas por dificuldade financeira, por qualquer das duas classificações, não altera a ordenação.",
    'I.4: saidas recalculadas e motivo pelo cadastro')
sub('Bonferroni sobre os valores do bootstrap.', 'Bonferroni sobre os valores do bootstrap, para as seis comparações contra benchmarks e referência não linear.', 'I.5: familia de Bonferroni')
sub('Precisão e revocação, Brier, calibração, tabela operacional com limiares fixados no treino e decisão líquida.',
    'Precisão e revocação, Brier, calibração logística sobre o posto do escore, tabela operacional com limiares fixados no treino e decisão líquida, nos dois horizontes.',
    'I.6: calibracao sobre o posto')
sub('Retirado como extensão prevista.', 'Mantido apenas como agenda, sem compromisso de execução no artigo.', 'indice agregado: coerente com a 6.2')

# quadro 3
sub('PCA 0,764; árvores 0,900; cobertura de juros 0,894', 'PCA 0,764; árvores 0,897; cobertura de juros 0,894', 'quadro 3: arvores')
sub('PCA 0,627; árvores 0,752', 'PCA 0,627; árvores 0,753', 'quadro 3: arvores')
sub('patrimônio: PCA 0,858, o melhor;', 'patrimônio: PCA 0,858, sem diferença distinguível;', 'quadro 3: empate, nao superioridade')
sub('PCA 0,857; PLS 0,856; árvores 0,858', 'PCA 0,857; PLS 0,856; árvores 0,857', 'quadro 3: arvores')

# secao 6
sub('em trinta firmas o primeiro documento trata do andamento do processo, e a sensibilidade correspondente está declarada como limitação',
    'em 22 firmas do painel o primeiro documento trata do andamento do processo, e a sensibilidade que antecipa em um ano a data dessas firmas foi estimada e não altera a conclusão central',
    'secao 6: 22 firmas e teste feito')
alvo = next((p for p in d.paragraphs if p.text.startswith('A terceira decorre da conferência bibliográfica')), None)
if alvo is not None:
    insert_after(alvo, 'Uma revisão integral do código e dos dados, feita depois desta rodada, corrigiu seis pontos, nenhum dos quais altera a conclusão central: '
                 'a identificação das saídas do painel usava um ano fixo, anterior ao fim do painel estendido; o sinal do segundo componente na verificação da soma de componentes não estava fixado entre partições; '
                 'a correção de Bonferroni considera seis comparações, e não sete; o modelo de Brito e Assaf Neto dentro da amostra não era winsorizado; a posição Ponzi passou a exigir exercícios de fato consecutivos; '
                 'e a calibração em probabilidade passou a ser feita sobre o posto do escore, o que tornou todos os modelos melhores que a previsão constante no escore de Brier. '
                 'A mesma revisão quantificou a sobreposição mecânica no horizonte de um ano, estimou a sensibilidade à datação da recuperação judicial e conferiu na fonte o que o texto afirma de cada obra citada.')
    log.append('OK     secao 6: paragrafo com as correcoes da revisao')
else:
    log.append('FALHOU secao 6: ancora nao encontrada')

d.save(DST)
print('\n'.join(log)); print(f"\n{sum(l.startswith('OK') for l in log)} aplicadas, {sum(l.startswith('FALHOU') for l in log)} falhas")
